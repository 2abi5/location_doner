"""Create reference.docx (Word styles used by pandoc): Times New Roman, 1.5 spacing, A4, 2.5 cm margins,
page numbers in the footer, continuous line numbers, black headings, booktabs-style tables.
Usage: ../.venv/bin/python make_reference_docx.py"""
import subprocess
import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

with open("reference.docx", "wb") as fh:
    fh.write(subprocess.run([pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
                            check=True, capture_output=True).stdout)
doc = Document("reference.docx")
BLACK = RGBColor(0, 0, 0)


def font(style, size, bold=None, italic=None, name="Times New Roman"):
    f = style.font
    f.name, f.size, f.color.rgb = name, Pt(size), BLACK
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]


def para(style, before=0, after=6, spacing=1.5, align=None):
    pf = style.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    if align is not None: pf.alignment = align


S = doc.styles
for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Bibliography"):
    if name in [s.name for s in S]:
        font(S[name], 12); para(S[name], after=6)
para(S["Compact"], after=2)
para(S["Bibliography"], after=4, spacing=1.15)
font(S["Bibliography"], 11)
S["Bibliography"].paragraph_format.left_indent = Cm(0.75)
S["Bibliography"].paragraph_format.first_line_indent = Cm(-0.75)
for name, size, bold, italic, before in (("Heading 1", 14, True, False, 18), ("Heading 2", 12, True, False, 12),
                                         ("Heading 3", 12, True, True, 10)):
    font(S[name], size, bold, italic); para(S[name], before=before, after=6, spacing=1.15)
font(S["Title"], 17, True); para(S["Title"], after=6, spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
font(S["Subtitle"], 11, False, True); para(S["Subtitle"], after=10, spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
for name in ("Author", "Date"):
    font(S[name], 11); para(S[name], after=4, spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
font(S["Abstract"], 11); para(S["Abstract"], after=6, spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
for name in ("Image Caption", "Table Caption"):
    font(S[name], 10, False, False); para(S[name], before=4, after=10, spacing=1.15)
from docx.enum.style import WD_STYLE_TYPE
if "Source Code" not in [x.name for x in S]:
    sc = S.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH); sc.base_style = S["Normal"]
font(S["Source Code"], 8.5, name="Courier New"); para(S["Source Code"], after=0, spacing=1.0)
font(S["Verbatim Char"], 8.5, name="Courier New")
font(S["Abstract Title"], 11, True)

# booktabs-like table style: rule above and below the table, thin rule under the header row
tbl = S["Table"].element
tblPr = tbl.find(qn("w:tblPr"))
if tblPr is None:
    tblPr = OxmlElement("w:tblPr"); tbl.append(tblPr)
for old in tblPr.findall(qn("w:tblBorders")):
    tblPr.remove(old)
borders = OxmlElement("w:tblBorders")
for side, sz in (("top", "12"), ("bottom", "12")):
    e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz); e.set(qn("w:color"), "000000")
    borders.append(e)
tblPr.append(borders)
for old in tbl.findall(qn("w:tblStylePr")):
    tbl.remove(old)
first = OxmlElement("w:tblStylePr"); first.set(qn("w:type"), "firstRow")
tcPr = OxmlElement("w:tcPr"); tcb = OxmlElement("w:tcBorders"); b = OxmlElement("w:bottom")
b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:color"), "000000")
tcb.append(b); tcPr.append(tcb); first.append(tcPr)
rpr = OxmlElement("w:rPr"); bold = OxmlElement("w:b"); rpr.append(bold); first.append(rpr)
tbl.append(first)
font(S["Table"], 10)

# page set-up, footer page number, continuous line numbering
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))
fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
fp.text = ""; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if kind:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)
ln = OxmlElement("w:lnNumType")
ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "283")
sec._sectPr.append(ln)
doc.save("reference.docx")
print("reference.docx written")
